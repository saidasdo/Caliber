"""Read `Equipment Performance - RCA{n}_{TAG}.xlsx`: Equipment Info, Condition History, Performance Summary."""

from dataclasses import dataclass, field

import openpyxl

from app.ingest.transforms.normalize import split_param_unit

EQUIPMENT_INFO_KEYS = {
    "Equipment Tag": "tag",
    "Equipment Name": "name",
    "Equipment Type": "eq_type_text",
    "Equipment Class": "eq_class",
    "Plant / Unit": "plant_unit_raw",
    "Discipline": "discipline",
    "Criticality": "criticality",
    "Design Life": "design_life",
    "Monitoring Method": "monitoring_method",
    "Linked RCA / AR No.": "linked_ar_no",
    "Failure Date": "failure_date_raw",
    "Dominant Failure Mode": "dominant_failure_mode",
}

PERFORMANCE_SUMMARY_KEYS = {
    "Monitoring Period (weeks)": "monitoring_period_weeks",
    "Total Downtime (hours)": "total_downtime_hours",
    "Period Hours": "period_hours",
    "Availability (%)": "availability_pct",
    "No. of Failures (period)": "failures_period",
    "MTBF (hours)": "mtbf_hours",
    "MTTR (hours)": "mttr_hours",
    "ALARM readings": "alarm_readings",
    "TRIP readings": "trip_readings",
    "NORMAL readings": "normal_readings",
    "PM Compliance (%)": "pm_compliance_pct",
    "Production Loss (ton)": "production_loss_ton",
    "Estimated Loss (k USD)": "estimated_loss_kusd",
}


@dataclass
class EquipmentPerformance:
    info: dict = field(default_factory=dict)
    limits: list = field(default_factory=list)          # [{parameter, unit, alarm_trip_raw}]
    condition_weekly: list = field(default_factory=list)  # [{week, date, parameter, unit, value, health_status, remark}]
    performance_summary: dict = field(default_factory=dict)


def read_equipment_performance_file(path) -> EquipmentPerformance:
    wb = openpyxl.load_workbook(path, data_only=True)
    result = EquipmentPerformance()

    ws = wb["Equipment Info"]
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
        key_a, val_a, key_c, val_c = (list(row) + [None] * 4)[:4]
        if key_a in EQUIPMENT_INFO_KEYS:
            result.info[EQUIPMENT_INFO_KEYS[key_a]] = val_a
        if key_c == "Parameter":
            continue
        if key_c and val_c is not None:
            name, unit = split_param_unit(key_c)
            result.limits.append({"parameter": name, "unit": unit, "alarm_trip_raw": val_c})

    ws = wb["Condition History"]
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    header = rows[0]
    param_cols = [(i, *split_param_unit(h)) for i, h in enumerate(header) if i not in (0, 1, 6, 7)]
    for row in rows[1:]:
        week, date = row[0], row[1]
        if week is None:
            continue
        health_status, remark = row[6], row[7]
        for idx, param_name, unit in param_cols:
            result.condition_weekly.append(
                {
                    "week": week,
                    "date": str(date),
                    "parameter": param_name,
                    "unit": unit,
                    "value": row[idx],
                    "health_status": health_status,
                    "remark": remark,
                }
            )

    ws = wb["Performance Summary"]
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
        key, value = row[0], row[1]
        if key in PERFORMANCE_SUMMARY_KEYS:
            result.performance_summary[PERFORMANCE_SUMMARY_KEYS[key]] = value

    return result
